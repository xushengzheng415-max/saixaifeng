from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"E:\Documents\sxf-football\docs\赛小蜂赛事球队与球员标准化体系（足球篮球通用）.docx")
PROTOTYPE_ROOT = Path(
    r"E:\Documents\saixiaofeng_football\赛小蜂足球UI\赛小蜂足球UI\原型图2.0\07-标准入驻链路"
)

GREEN = "0B6B3A"
DARK_GREEN = "083F28"
MID_GREEN = "2E8B57"
LIGHT_GREEN = "EAF6EF"
PALE_GREEN = "F5FBF7"
GOLD = "C89B3C"
ORANGE = "D97706"
RED = "B42318"
INK = "1F2937"
MUTED = "667085"
BORDER = "B9D8C5"
WHITE = "FFFFFF"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            width = widths_dxa[min(i, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=MUTED)


def set_paragraph_border_bottom(paragraph, color=GREEN, size="14", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    style_map = {
        "Title": (28, DARK_GREEN, 0, 8),
        "Subtitle": (13, MUTED, 0, 16),
        "Heading 1": (16, GREEN, 18, 10),
        "Heading 2": (13, GREEN, 14, 7),
        "Heading 3": (11.5, DARK_GREEN, 10, 5),
    }
    for name, (size, color, before, after) in style_map.items():
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("赛小蜂赛事标准体系  |  足球 / 篮球通用")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    set_paragraph_border_bottom(p, color=BORDER, size="6", space="3")

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_body(doc, text, bold_lead=None, color=INK, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, title, text, tone="green"):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9120], indent=120)
    cell = table.cell(0, 0)
    fill = LIGHT_GREEN if tone == "green" else ("FFF7E8" if tone == "orange" else "F5F5F5")
    accent = GREEN if tone == "green" else (ORANGE if tone == "orange" else MUTED)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    set_run_font(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GREEN, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, indent=120)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2, bold=True, color=DARK_GREEN)
    for row_idx, row_values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_GREEN)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, image_path, caption, width=6.25):
    path = Path(image_path)
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    r = cp.add_run(caption)
    set_run_font(r, size=8.5, italic=True, color=MUTED)


def add_section_break(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)
    section = doc.sections[-1]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True


def build_document():
    doc = Document()
    setup_styles(doc)
    doc.core_properties.title = "赛小蜂赛事球队与球员标准化体系（足球/篮球通用）"
    doc.core_properties.subject = "赛事创建、球队入驻、球员实名、名单与比赛执行的跨项目产品标准"
    doc.core_properties.author = "赛小蜂产品团队"
    doc.core_properties.keywords = "赛小蜂, 足球, 篮球, 赛事SaaS, 球队, 球员, 实名认证"

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("赛小蜂赛事产品标准")
    set_run_font(r, size=11, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("球队与球员标准化体系")
    set_run_font(r, size=30, bold=True, color=DARK_GREEN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("足球 / 篮球通用版")
    set_run_font(r, size=17, bold=True, color=GREEN)
    add_callout(
        doc,
        "标准目标",
        "让球队必须在赛小蜂小程序完成业务落位，同时允许主办方代录、球队维护和家长服务号补充自然协作；将长期球队、球员、赛事名单、比赛事件和球队数据包严格分层。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("版本 1.1  |  已确认基线  |  2026-07-29")
    set_run_font(r, size=10, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("适用于赛小蜂足球、赛小蜂篮球及后续团队球类赛事产品")
    set_run_font(r, size=9.5, color=MUTED)

    doc.add_page_break()

    # Executive summary / contents
    doc.add_heading("使用说明", level=1)
    add_body(
        doc,
        "本文件是产品、原型、研发和运营共同遵守的标准基线。运动项目之间共用赛事、组别、球队、球员、人员关系、资格、名单和权限框架；只有阵容规则、比赛事件和统计口径进入项目适配层。",
    )
    add_callout(
        doc,
        "一句话标准",
        "主办方创建赛事后可邀请球队，球队也可在小程序搜索赛事申请；无论从哪个入口进入，球队都必须在小程序创建或认领长期球队并确认参赛。家长、实名认证和形象照按赛事要求启用，不与专业版强绑定。",
    )
    doc.add_heading("目录", level=2)
    contents = [
        "1. 产品目标与非目标",
        "2. 跨项目赛事总框架",
        "3. 账号、关系与端口职责",
        "4. 从赛事创建到球队入驻",
        "5. 球员最小建档与家长补充",
        "6. 身份证实名与标准形象照",
        "7. 正式名单、单场阵容与比赛快照",
        "8. 简易版与专业版",
        "9. 足球与篮球适配矩阵",
        "10. 状态流、异常与权限",
        "11. 产品页面标准与验收",
        "12. 数据与实施建议",
    ]
    add_numbers(doc, contents)

    # 1
    doc.add_heading("1. 产品目标与非目标", level=1)
    doc.add_heading("1.1 产品目标", level=2)
    add_bullets(
        doc,
        [
            "主办方可以代录球队和球员，也可以让球队或家长补充，不被单一资料流程限制。",
            "所有参赛球队必须进入赛小蜂小程序创建或认领球队并确认参赛，形成长期业务关系。",
            "球队无需在表格、PC和小程序之间重复录入同一批球员。",
            "家长通过统一链接只补充自己的孩子，不接触整队隐私。",
            "实名认证由主办方独立设置；启用后一次实名、经授权跨赛事复用。",
            "足球与篮球共享同一业务底座，不复制两套相互冲突的身份和数据系统。",
        ],
    )
    doc.add_heading("1.2 明确不做", level=2)
    add_bullets(
        doc,
        [
            "不要求主办方为了建队填写省市、地址、教练和球员等非必要信息。",
            "不要求教练上传球员照片、身份证、学籍或保险材料。",
            "不把身份证实名、学籍证明或活体刷脸作为专业版固定要求；实名由主办方独立开启。",
            "不让主办方在开赛前选择“主办方录入/球队协作/家长协作”三种互斥模式。",
            "不让球员从公开入口搜索同名档案并直接认领。",
            "不把可变化的地区、年龄组、球队和号码编码成自然人的永久身份。",
        ],
    )

    # 2
    doc.add_heading("2. 跨项目赛事总框架", level=1)
    add_body(
        doc,
        "赛事系统使用“稳定主体 + 赛事关系 + 快照”的三层模型。球队和球员是长期资产；报名、名单、阵容和事件属于具体赛事或比赛；历史快照不得被日常资料修改覆盖。",
    )
    add_table(
        doc,
        ["层级", "核心对象", "标准含义", "是否长期复用"],
        [
            ["平台层", "User / Person", "自然人账号与人物身份", "是"],
            ["机构层", "Organization", "俱乐部、学校、赛事公司等工作空间", "是"],
            ["球队层", "Team", "名称、队徽及长期球队身份", "是"],
            ["球员层", "PlayerProfile", "球员自然人档案与实名状态", "是"],
            ["赛事层", "Tournament", "一届独立赛事", "否"],
            ["组别层", "Division", "U9、公开组等；版本与收费单位", "否"],
            ["报名层", "TournamentRegistration", "球队参加赛事/组别的关系", "否"],
            ["名单层", "RosterSnapshot", "审核通过的正式参赛名单快照", "否"],
            ["比赛层", "Match / MatchLineup", "具体场次、首发/上场阵容与替补", "否"],
            ["事件层", "MatchEvent", "得分、犯规、换人等可追溯事件", "否"],
        ],
        [1200, 1880, 4200, 1840],
    )
    doc.add_heading("2.1 编号原则", level=2)
    add_callout(
        doc,
        "编号跟赛事走，身份不跟编号走",
        "数据库内部使用不可变系统ID。对外可读的参赛编号由赛事自动生成，只用于当届赛事的搜索、表格、名单和客服定位。球员转队或球队年龄组变化时，稳定身份不改变。",
    )
    add_table(
        doc,
        ["编号", "示例", "用途", "用户是否填写"],
        [
            ["球队系统ID", "内部UUID", "长期识别同一球队", "否"],
            ["球员系统ID", "内部UUID", "长期识别同一自然人", "否"],
            ["赛事参赛编号", "HNYC-U9-009", "本届赛事球队定位", "否，系统生成"],
            ["正式名单编号", "U9-A-012", "本届名单与纸质材料定位", "否，系统生成"],
            ["球衣号码", "8 / 10 / 23", "球队、赛事或比赛关系字段", "由球队选择"],
        ],
        [1700, 1900, 3300, 2460],
    )

    # 3
    doc.add_heading("3. 账号、关系与端口职责", level=1)
    add_body(
        doc,
        "身份不采用永久单角色。权限来自自然人与球队、机构、赛事、组别或场次之间的关系。业务入口目前以球队和裁判为主，球员/监护人通过球队发起的专属链接进入。",
    )
    add_table(
        doc,
        ["角色", "主要端口", "核心职责", "禁止越权"],
        [
            ["主办方", "PC + 主办方小程序", "创建赛事、开放入驻、配置组别、处理异常、确认名单", "不进入球队私有经营数据"],
            ["球队负责人/教练", "赛小蜂小程序", "创建/认领长期球队、申请/确认参赛、按权限维护球员与名单", "不直接修改实名结果"],
            ["家长/监护人", "服务号页面/H5", "按赛事要求核对孩子、实名、完成照片与授权", "不能查看整队资料"],
            ["裁判", "裁判手机端/小程序", "接受任务、核验阵容、执行比赛记录", "不能建立或修改球员长期档案"],
            ["观众", "公开H5", "查看公开赛程、比分、排名与授权展示内容", "不能接触身份证和监护人信息"],
        ],
        [1500, 1700, 3800, 2360],
    )
    doc.add_heading("3.1 端口落位标准", level=2)
    add_bullets(
        doc,
        [
            "球队业务最终必须落在赛小蜂小程序：微信授权、创建/认领球队、加入赛事和赛事工作台均在小程序完成。",
            "主办方邀请链接、小程序码或服务号菜单都可以引导球队打开小程序，但不能替代小程序业务落位。",
            "家长可以只在服务号页面/H5完成自己的孩子资料，不要求进入球队小程序。",
            "球员/监护人不从公共入口搜索同名档案，而从球队发出的本队专属链接进入。",
            "首次绑定后，微信身份用于登录与通知；手机号用于验证；权限来自已审核关系。",
        ],
    )

    # 4
    doc.add_heading("4. 从赛事创建到球队入驻", level=1)
    doc.add_heading("4.1 标准主路径", level=2)
    add_numbers(
        doc,
        [
            "主办方创建赛事空间并保存名称、时间、地点等基础信息。",
            "入口一：主办方发送赛事小程序链接或小程序码，球队在小程序创建/认领球队并确认加入。",
            "入口二：球队先在小程序创建长期球队，再搜索公开赛事并提交参赛申请。",
            "两种入口可以同时存在，不是互斥模式；主办方可以关闭公开搜索申请，但邀请入口始终存在。",
            "球队只填写球队名称和队徽；主办方也可提前建立球队壳子供负责人认领。",
            "主办方审核加入关系后，系统生成本届赛事参赛编号。",
            "球队进入本队赛事工作台，接收赛程、比分、排名、通知和名单任务。",
            "简易版不要求提交赛事参赛球员；专业版再从球队球员库选择本届参赛名单。",
        ],
    )
    doc.add_heading("4.2 球队最小建档", level=2)
    add_table(
        doc,
        ["字段", "要求", "说明"],
        [
            ["球队名称", "必填", "主办方或球队负责人输入"],
            ["球队队徽", "建议上传", "可拍摄、相册选择或暂用默认队徽"],
            ["赛事参赛编号", "系统生成", "仅属于当届赛事，不要求记忆"],
            ["负责人手机号", "认领时验证", "不是微信号；可与微信绑定手机号不同"],
            ["其他资料", "入驻后按需完善", "不阻断球队建档"],
        ],
        [1800, 1800, 5760],
    )
    add_callout(
        doc,
        "主办方减负规则",
        "主办方只处理邀请、搜索申请、认领审核和异常。球队必须进小程序，但资料可由主办方代录、球队维护或家长按需补充。",
    )

    add_figure(
        doc,
        PROTOTYPE_ROOT / "01-主办方PC" / "05-球队加入赛事设置.png",
        "场景示例：主办方邀请与球队主动申请两个并行入口",
    )

    # 5
    doc.add_heading("5. 球员最小建档与家长补充", level=1)
    doc.add_heading("5.1 教练只建立球员壳子", level=2)
    add_table(
        doc,
        ["字段", "要求", "系统用途"],
        [
            ["球员姓名", "必填", "创建待补充档案与后续核对"],
            ["出生年月", "年龄组赛事按需", "计算比赛日年龄并做组别初筛"],
            ["监护人手机号", "开启家长补充时必填", "匹配家长、发送提醒、避免公开名单"],
        ],
        [2000, 1800, 5560],
    )
    add_body(
        doc,
        "主办方代录且不启用家长补充时，监护人手机号可以为空。年龄不作为固定字段保存；系统根据出生信息与赛事日期动态计算比赛日年龄。赛事规则确需精确到日且主办方开启实名时，再由身份证OCR补齐完整出生日期。",
    )
    doc.add_heading("5.2 两种家长路径", level=2)
    add_table(
        doc,
        ["路径", "适用场景", "家长动作", "球队动作"],
        [
            ["教练先建壳子", "教练已有姓名与联系方式", "验证手机号后匹配孩子并补充", "确认三字段并发统一链接"],
            ["家长直接登记", "教练尚未建立完整名单", "从群链接登记孩子并提交", "确认该孩子属于本队"],
        ],
        [1800, 2500, 2900, 2160],
    )
    add_bullets(
        doc,
        [
            "家长补充是赛事可选能力；未开启时不生成家长待办或链接。",
            "开启后，统一家长链接可以直接发到家长群，不要求教练逐个转发。",
            "家长验证手机号后只看到自己的孩子，不能浏览整队名单。",
            "同一手机号允许关联多个孩子；新增孩子必须由球队负责人确认。",
            "球员档案可以在未上传照片时先存在，状态为“待家长补充”。",
        ],
    )
    add_figure(
        doc,
        PROTOTYPE_ROOT / "02-球队小程序" / "05-快速添加球员壳子.png",
        "场景示例：教练只填写姓名、出生年月和监护人手机号",
    )

    # 6
    doc.add_heading("6. 身份证实名与标准形象照", level=1)
    doc.add_heading("6.1 可选实名认证标准", level=2)
    add_callout(
        doc,
        "当前已确认标准",
        "实名认证不是专业版默认必选项。只有主办方在赛事或组别中开启后，家长才上传身份证人像面并完成姓名与身份证号核验；不要求学籍证明，不要求球员活体刷脸。",
    )
    add_numbers(
        doc,
        [
            "家长上传或拍摄身份证人像面。",
            "OCR读取姓名、身份证号、出生日期、性别和证件头像。",
            "家长核对脱敏信息并确认。",
            "系统进行姓名与身份证号二要素核验。",
            "通过后生成平台级已实名球员档案；后续赛事经授权复用。",
            "年龄不符、重复报名、识别失败等进入异常队列。",
        ],
    )
    add_table(
        doc,
        ["状态", "含义", "下一动作"],
        [
            ["待实名", "尚未上传身份证", "提醒家长"],
            ["识别中", "OCR或核验处理中", "等待系统结果"],
            ["已实名", "姓名与身份证号匹配", "继续形象照或球队确认"],
            ["实名失败", "信息不匹配或无法核验", "家长重拍/更正"],
            ["资格异常", "年龄、重复报名等规则异常", "主办方处理"],
        ],
        [1700, 4300, 3360],
    )
    doc.add_heading("6.2 标准球员形象照", level=2)
    add_bullets(
        doc,
        [
            "主入口为“立即拍摄标准形象照（推荐）”，次入口为“从相册选择”。",
            "建议穿本俱乐部比赛球衣，正面站立，双臂交叉于胸前，拍摄到腰部。",
            "相机展示半透明标准姿势取景框，并提示距离、光线、清晰度、头部和手臂完整性。",
            "平台自动完成人像分割，一次拍摄生成透明背景形象照和头肩头像两种资产。",
            "家长可拖动、缩放头像裁切框并确认圆形/方形头像，无需重复拍摄。",
            "系统保留原图、透明人物图、头像裁切参数和标准预览。",
            "球衣为强建议，不作为提交阻断；质量不合格时给出可理解的重拍原因。",
        ],
    )
    doc.add_heading("6.3 敏感数据边界", level=2)
    add_bullets(
        doc,
        [
            "身份证原图、实名结果和公开球员形象照分开存储、分开授权。",
            "球队和裁判默认只看实名状态、年龄资格和脱敏字段，不查看身份证原图。",
            "身份证原图需要加密、访问审计和明确保留期限；正式上线前完成未成年人隐私合规评审。",
            "实名结果与球员系统ID绑定，不因转队或参加新赛事而重复建立自然人身份。",
        ],
    )

    # 7
    doc.add_heading("7. 正式名单、单场阵容与比赛快照", level=1)
    add_body(
        doc,
        "球员基础档案、球队日常成员、赛事正式名单和单场阵容是四种不同数据。任何页面不得将它们混为同一个可随意覆盖的列表。",
    )
    add_callout(
        doc,
        "三层名单不得混用",
        "球队球员库是长期资产；赛事参赛名单是从球队库选择并绑定具体赛事/组别的子集；单场比赛名单只能从已审核赛事名单中选择。球队加入赛事时不得自动提交全队球员。",
    )
    add_table(
        doc,
        ["对象", "形成时间", "负责人", "是否可覆盖历史"],
        [
            ["PlayerProfile", "首次实名建档", "家长/平台", "否"],
            ["TeamMembership", "加入球队时", "球队负责人", "保留起止时间"],
            ["RosterSnapshot", "赛事报名审核通过时", "球队提交、主办方确认", "否"],
            ["MatchLineup", "每场赛前", "球队提交、裁判核验", "否"],
            ["MatchEvent", "比赛执行中", "授权记录员/裁判", "只能审计更正"],
        ],
        [1900, 2100, 3000, 2360],
    )
    doc.add_heading("7.1 名单形成标准", level=2)
    add_numbers(
        doc,
        [
            "球员完成本组已开启的必要资料；未开启实名认证时不得产生实名待办。",
            "球队负责人从球队球员库选择本届、本组别参赛球员，不自动全选。",
            "系统自动执行年龄、重复和完整度校验。",
            "主办方批量确认正常球员，只处理异常。",
            "生成不可被日常资料修改覆盖的RosterSnapshot。",
            "每场比赛从RosterSnapshot选择首发、替补或上场名单。",
        ],
    )

    # 8
    doc.add_heading("8. 简易版与专业版", level=1)
    add_table(
        doc,
        ["能力", "简易版", "专业版"],
        [
            ["收费单位", "免费", "按竞赛组别开通"],
            ["球队建档", "队名与队徽", "队名与队徽"],
            ["赛事参赛球员", "不要求提交", "从球队球员库选择并形成正式名单"],
            ["实名认证", "主办方独立设置", "主办方独立设置，不因专业版自动开启"],
            ["阵容", "不提供结构化球员阵容", "从正式名单提交首发/上场与替补"],
            ["比赛事件", "可选人工文字", "必须关联正式球员"],
            ["个人统计", "不生成赛事公开球员榜", "自动生成赛事官方统计"],
            ["赛后资料", "基础纸质资料电子归档", "系统自动生成正式电子记录"],
        ],
        [1900, 3680, 3780],
    )
    add_callout(
        doc,
        "锁定原则",
        "组别可以在确认竞赛方案前从简易版升级专业版；确认竞赛方案并进入正式比赛管理后，当前版本永久锁定，不再升级或降级。",
        tone="orange",
    )
    doc.add_heading("8.1 球队赛事数据包", level=2)
    add_body(
        doc,
        "当主办方未购买组别专业版时，单支球队仍可购买球队赛事数据包，把简易版裁判记录的本队球员事件沉淀为球队自己的长期数据。该权益不改变赛事版本，不生成赛事公开球员榜。",
    )
    add_table(
        doc,
        ["场景", "裁判录入方式", "数据结果"],
        [
            ["球队已购数据包", "选择事件所属球队后，直接从该队球员库选择球员", "绑定稳定球员ID，进入个人档案、累计统计和导出接口"],
            ["球队未购数据包", "手工输入球衣号码或姓名", "只保留为单场文字记录，不累计个人档案"],
            ["同场一方购买", "购买方选球员，未购买方手工输入", "只为购买方沉淀长期球员数据"],
            ["录入异常", "球员未找到或临时新增", "进入赛后待确认，正常事件不重复匹配"],
        ],
        [1900, 3900, 3560],
    )
    add_callout(
        doc,
        "数据出口",
        "球队数据包形成的个人比赛记录和累计数据保存在球队小程序，并支持导出；后续为球队自有教务系统保留标准接口。",
    )

    # 9
    doc.add_heading("9. 足球与篮球适配矩阵", level=1)
    add_body(
        doc,
        "以下能力使用同一底座：赛事、组别、球队小程序落位、球员档案、可选实名、监护人、三层名单、球队数据包、权限、通知、异常、支付权益和历史快照。差异只进入运动项目适配层。",
    )
    add_table(
        doc,
        ["模块", "足球适配", "篮球适配"],
        [
            ["位置", "门将、后卫、中场、前锋", "后卫、前锋、中锋或自定义位置"],
            ["单场首发", "按赛制选择首发人数及替补", "通常首发5人，其余替补"],
            ["球衣号码", "赛事/球队关系字段", "赛事/球队关系字段"],
            ["核心得分事件", "进球、点球", "罚球、两分、三分"],
            ["纪律/犯规", "黄牌、红牌", "个人犯规、技术犯规、违体犯规"],
            ["换人", "换上/换下事件", "频繁换人、场上五人阵容"],
            ["时钟", "上下半场或节次", "节次、比赛时钟、24秒等按配置"],
            ["球队统计", "射门、射正、角球等", "篮板、助攻、抢断、盖帽、失误等"],
            ["比赛终端", "裁判/比赛记录员轻量端", "计分员高压计分台"],
            ["简易版球队数据包", "购买方直接选择球员记录进球/牌/换人", "购买方直接选择球员记录得分/犯规/技术统计"],
        ],
        [1700, 3820, 3840],
    )
    doc.add_heading("9.1 篮球复用时必须保留", level=2)
    add_bullets(
        doc,
        [
            "球队和球员建立流程完全复用，不重新设计第二套实名体系。",
            "篮球位置、首发五人、比赛阵容和技术统计是适配字段，不写进通用球员基础档案。",
            "篮球计分台必须支持大触控、+1/+2/+3、时钟、节次、犯规、暂停、换人、撤销和审计。",
            "技术统计事件必须关联RosterSnapshot中的正式球员，避免同名文本产生重复球员。",
            "篮球比分、犯规、暂停和阵容需要离线重试、冲突解决和事件幂等。",
        ],
    )

    # 10
    doc.add_heading("10. 状态流、异常与权限", level=1)
    doc.add_heading("10.1 球队状态", level=2)
    add_table(
        doc,
        ["状态", "进入条件", "退出条件"],
        [
            ["待认领", "主办方建立球队壳子", "负责人认领通过"],
            ["待审核加入", "球队搜索赛事并提交申请", "主办方通过或驳回"],
            ["认领审核中", "联系人信息不一致", "主办方通过或驳回"],
            ["已加入赛事", "负责人关系与参赛关系建立", "进入本队赛事工作台"],
            ["待选择参赛球员", "专业版组别要求正式名单", "球队从球员库选择并提交"],
            ["待完善球员", "本组已启用资料要求且球员未完成", "球员资料完成"],
            ["待球队确认", "球员已补充", "球队确认提交"],
            ["正式名单已形成", "主办方确认", "只允许名单变更流程"],
        ],
        [1900, 3700, 3760],
    )
    doc.add_heading("10.2 球员状态", level=2)
    add_table(
        doc,
        ["状态", "责任人", "核心动作"],
        [
            ["待家长补充", "家长", "打开群链接并验证手机号"],
            ["待实名", "家长", "仅在主办方开启实名后上传身份证人像面"],
            ["实名失败", "家长", "重新拍摄或更正"],
            ["待形象照", "家长", "仅在主办方开启照片要求后拍摄并裁切头像"],
            ["待球队确认", "球队负责人", "确认属于本队"],
            ["资料完成", "系统", "进入正常名单队列"],
            ["资格异常", "主办方", "处理年龄、重复或争议"],
        ],
        [2200, 2100, 5060],
    )
    doc.add_heading("10.3 球队赛事数据包状态", level=2)
    add_table(
        doc,
        ["状态", "现场录入", "赛后结果"],
        [
            ["未开通", "手工输入球衣号/姓名", "单场文字记录，不累计个人档案"],
            ["已开通", "直接选择球队球员", "绑定球员ID并累计到球队数据"],
            ["待确认异常", "未找到球员或临时新增", "球队赛后确认后归档"],
        ],
        [2200, 3300, 3860],
    )
    doc.add_heading("10.4 异常优先", level=2)
    add_bullets(
        doc,
        [
            "正常资料不要求主办方逐条点击，可以按球队或组别批量确认。",
            "手机号不一致、重复认领、疑似同人、年龄不符和实名失败进入独立异常队列。",
            "任何退回必须给出原因；原始提交、修改前后值、操作者和时间均保留。",
            "裁判核验只影响单场阵容状态，不反向修改球员长期实名档案。",
        ],
    )

    # 11
    doc.add_heading("11. 产品页面标准与验收", level=1)
    add_table(
        doc,
        ["端口", "必须页面", "核心验收"],
        [
            ["主办方PC", "球队加入设置、邀请/申请审核、快速建队、批量导入、实名异常、正式名单", "不选择三种模式；异常可定位"],
            ["球队小程序", "赛事中心、邀请加入、搜索申请、认领/建队、球队级工作台、名单选择、数据包", "球队必须落位；球队库与赛事名单分层"],
            ["家长服务页", "匹配孩子、可选实名、拍照指引、头像裁切、提交结果", "只见自己的孩子；未开启则不出现"],
            ["裁判手机端", "任务、阵容核验、混合权益事件录入、提交", "购买方选球员；未购买方手工输入"],
            ["公开H5", "赛程、比分、排名、授权球员展示", "不暴露敏感数据"],
        ],
        [1500, 4250, 3610],
    )
    doc.add_heading("11.1 关键体验指标", level=2)
    add_bullets(
        doc,
        [
            "主办方生成邀请小程序码并设置是否接受搜索申请：不超过2分钟。",
            "球队创建：队名与队徽两项，通常1分钟内完成。",
            "教练建立单个球员壳子：通常20秒内完成。",
            "家长在服务号完成已开启的实名与照片要求，不依赖PC或球队小程序。",
            "一次拍摄同时生成形象照与可手动裁切头像。",
            "裁判为已购数据包球队记录事件时，最多三步完成：选事件、选球员、确认。",
            "主办方工作台默认展示完成度与异常，而非全量资料录入表。",
            "任何页面不得要求用户再次填写系统已经识别或已经确认的数据。",
        ],
    )
    doc.add_heading("11.2 原型验收清单", level=2)
    add_bullets(
        doc,
        [
            "跨页面的赛事名称、球队名称、参赛编号、人数和状态一致。",
            "简易版与专业版边界一致，比赛管理阶段不再出现升级营销。",
            "球队加入入口只有主办方邀请和小程序搜索申请，不出现三种协作模式。",
            "球队球员库、赛事参赛名单和单场比赛名单没有自动混用。",
            "实名认证、家长补充和标准照片关闭时，不产生相应待办。",
            "同场一队购买数据包时，裁判界面按球队切换选择球员/手工输入。",
            "身份证原图不出现在球队列表、裁判页和公开页。",
            "球员基础档案不混入位置、号码、阵容、保险或单场统计。",
            "足球与篮球适配字段不会污染通用实体。",
            "空状态、失败、退回、重复、离线和冲突均有明确处理入口。",
        ],
    )

    # 12
    doc.add_heading("12. 数据与实施建议", level=1)
    doc.add_heading("12.1 最小实体建议", level=2)
    add_table(
        doc,
        ["实体", "关键字段"],
        [
            ["User", "id, unionId/openId, verifiedPhones"],
            ["OrganizationMembership", "userId, organizationId, roles, permissions"],
            ["Team", "id, name, logo, organizationId, status"],
            ["PlayerProfile", "id, name, birthDate/birthMonth, identityStatus, portraitAssets"],
            ["GuardianRelation", "playerId, userId, phone, relation, consentStatus"],
            ["TeamMembership", "teamId, playerId, startAt, endAt, status"],
            ["TournamentRegistration", "tournamentId, divisionId, teamId, registrationCode, status"],
            ["RosterSnapshot", "registrationId, playerIds, version, approvedAt"],
            ["MatchLineup", "matchId, rosterPlayerIds, starter/bench/active status"],
            ["MatchEvent", "matchId, teamId, playerId, type, period, clock, operatorId, revision"],
            ["DivisionEntitlement", "divisionId, professionalStatus, lockedAt"],
            ["TeamTournamentDataPack", "teamId, tournamentId/divisionId, status, validFrom, validTo"],
            ["TeamPlayerEventArchive", "teamId, playerId, matchEventId, source, archivedAt"],
        ],
        [2600, 6760],
    )
    doc.add_heading("12.2 兼容与迁移", level=2)
    add_bullets(
        doc,
        [
            "现有球队编号和球员编号可作为兼容业务字段保留，不继续承担永久身份主键。",
            "历史表格导入后先生成待确认记录，不直接覆盖已有球队或球员。",
            "同名球队和同名球员通过联系人、出生信息、历史关系与人工确认处理，不仅靠名称自动合并。",
            "足球和篮球应共享通用身份与入驻服务；比赛事件、阵容和统计使用sportType适配。",
            "简易版手工文字事件与已购数据包的结构化球员事件分开保存，不能通过姓名自动合并。",
            "组别专业版权益与单支球队赛事数据包使用不同权益实体和订单归属。",
            "当前数据库模型迁移需要单独方案，不能一次性破坏已有users、teams、players和赛事数据。",
        ],
    )
    doc.add_heading("12.3 上线顺序", level=2)
    add_numbers(
        doc,
        [
            "先上线球队小程序赛事中心、邀请加入、搜索申请和长期球队创建/认领。",
            "再上线主办方代录/批量导入、球队球员库和三层名单分离。",
            "上线专业版RosterSnapshot与足球比赛执行。",
            "上线简易版球队赛事数据包和按球队权益切换的裁判事件录入。",
            "按需接入家长服务号补充、身份证OCR、二要素实名和敏感数据权限。",
            "上线一次拍摄生成形象照与头像裁切。",
            "篮球复用通用底座，仅新增篮球阵容、计分、统计和数据包事件适配。",
        ],
    )
    add_callout(
        doc,
        "最终产品判断标准",
        "不是把录入工作从主办方转嫁给球队或家长，而是让平台承担识别、匹配、状态和复用；每个角色只确认自己最了解的信息。",
    )

    # Appendix
    doc.add_page_break()
    doc.add_heading("附录A：标准链路速查", level=1)
    add_table(
        doc,
        ["阶段", "主办方", "球队负责人", "家长", "平台"],
        [
            ["赛事创建", "创建并设置邀请/搜索申请", "—", "—", "生成小程序入口"],
            ["球队加入", "发邀请/审核申请", "小程序创建或认领并确认", "—", "建立长期球队与赛事关系"],
            ["球队球员库", "代录/批量导入", "按权限维护", "—", "长期保存，不自动成为赛事名单"],
            ["资料补充", "按需开启要求", "发服务号链接", "完成已开启项目", "OCR、人像分割与头像裁切"],
            ["赛事名单", "处理异常并确认", "从球队库选择本届球员", "必要时更正", "生成名单快照"],
            ["比赛执行", "配置与复核", "提交阵容", "查看通知", "生成比赛快照与统计"],
            ["简易版数据包", "无需购买专业版", "单队购买并查看沉淀数据", "—", "购买方选球员并累计；未购买方文字记录"],
        ],
        [1400, 2100, 2100, 1900, 1860],
        font_size=8.6,
    )
    doc.add_heading("附录B：当前已确认产品决策", level=1)
    add_bullets(
        doc,
        [
            "球队建档只要求球队名称和队徽。",
            "赛事参赛编号由系统自动生成并跟随具体赛事。",
            "球队加入赛事只有主办方邀请和小程序搜索申请两个并行入口。",
            "所有球队必须在赛小蜂小程序创建或认领球队并确认参赛。",
            "球队是长期资产；球队球员库、赛事参赛名单和单场比赛名单分层。",
            "主办方可代录/导入球员；球队可按权限维护；家长补充通过服务号按需开启。",
            "实名认证不与专业版强绑定；启用时只上传身份证人像面，不要求学籍或活体刷脸。",
            "一次拍摄生成透明形象照和可由家长拖动缩放的头像。",
            "简易版不要求赛事参赛球员；单支球队可购买数据包沉淀本队球员事件。",
            "数据包购买方裁判现场直接选球员，未购买方手工输入且不累计个人档案。",
            "足球与篮球共用赛事、球队、球员与实名底座。",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
