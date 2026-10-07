"""Make a generic Word registration form with native picture controls."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import sys
from docx import Document
from docx.table import _Cell
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'web-admin-vue/public/templates'
doc=Document(Path(sys.argv[1]))
for part in [doc]+[s.header for s in doc.sections]+[s.footer for s in doc.sections]:
    for p in part.paragraphs:
        if '立洋杯' in p.text:
            if '我自愿' in p.text: p.text=p.text.replace('2026年开封市第六届“立洋杯”足球联赛','本次足球赛事')
            else: p.text='足球赛事球队报名表'
        elif p.text.strip()=='足球邀请赛报名表': p.text='赛事名称：____________________    竞赛组别：____________________'
        if p.text.strip()=='附件2：':p.text='附件：'
doc.core_properties.title='足球赛事球队报名表'
doc.core_properties.subject='通用报名表；图片控件与可选人员字段'
doc.core_properties.author='赛小蜂足球'
doc.core_properties.last_modified_by='赛小蜂足球'
font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',36)
def placeholder(label):
    im=Image.new('RGB',(500,600),'white'); d=ImageDraw.Draw(im)
    d.rounded_rectangle((30,100,470,500),radius=14,fill='#F5F8F6',outline='#9CB6A6',width=3)
    d.line((220,245,280,245),fill='#087C47',width=6);d.line((250,215,250,275),fill='#087C47',width=6)
    d.text((250,340),label,font=font,fill='#30483A',anchor='mm')
    info=PngImagePlugin.PngInfo();info.add_text('sxf','SXF_UPLOAD_PLACEHOLDER_v1')
    stream=BytesIO();im.save(stream,format='PNG',pnginfo=info);return stream.getvalue()
serial=90000
def picture(cell,tag,label):
    global serial
    cell.text='';p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
    run=p.add_run();run.add_picture(BytesIO(placeholder(label)),width=Cm(2.4),height=Cm(2.88))
    sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    for name,value in [('alias',label),('tag',tag),('id',str(serial))]:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),value);pr.append(e)
    serial+=1;pr.append(OxmlElement('w:picture'));pr.append(OxmlElement('w:showingPlcHdr'))
    content=OxmlElement('w:sdtContent');p._p.replace(run._r,sdt);content.append(run._r);sdt.append(pr);sdt.append(content)
table=doc.tables[1]
for block,index in enumerate([19,14,9,4]):
    for col,tc in enumerate(table.rows[index+1]._tr.findall(qn('w:tc'))): picture(_Cell(tc,table),f'sxf-photo-{block}-{col}','选择照片')
    reference=table.rows[index+4]._tr
    for field,label in [('birthDate','出生日期：'),('position','位置：')]:
        row=deepcopy(reference)
        for tc in row.findall(qn('w:tc')):
            cell=_Cell(tc,table);cell.text=label
            if field=='position':
                p=cell.paragraphs[0]._p;sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');tag=OxmlElement('w:tag');tag.set(qn('w:val'),f'sxf-position-{serial}');pr.append(tag);serial+=1
                dropdown=OxmlElement('w:dropDownList')
                for name in ['请选择','门将','后卫','中场','前锋','其他']:
                    item=OxmlElement('w:listItem');item.set(qn('w:displayText'),name);item.set(qn('w:value'),name);dropdown.append(item)
                pr.append(dropdown);content=OxmlElement('w:sdtContent');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='请选择';r.append(t);content.append(r);sdt.append(pr);sdt.append(content);p.append(sdt)
        reference.addnext(row);reference=row
# Adding rows changes indices; the original row objects above remain valid only when processed backwards.
# The base source has merged header columns, but each person photo cell is a distinct tc.
logo=doc.tables[0].rows[0]._tr.findall(qn('w:tc'))[-1]
picture(_Cell(logo,doc.tables[0]),'sxf-logo','选择队徽')
for table in doc.tables:
    for row in table.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for height in list(pr.findall(qn('w:trHeight'))):pr.remove(height)
        texts=[_Cell(tc,table).text.strip() for tc in row._tr.findall(qn('w:tc'))]
        group_row=any(text.startswith(('姓名：','身份证：','球服号码：','出生日期：','位置：')) or text in ['运动员','主教练','队医','领队(联络员)'] for text in texts) or bool(row._tr.findall('.//'+qn('w:picture')))
        last_row=any(text.startswith('位置：') for text in texts)
        for tc in row._tr.findall(qn('w:tc')):
            for p in _Cell(tc,table).paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=group_row and not last_row
                for run in p.runs:run.font.size=Pt(10)
doc.save(OUT/'team-registration-form-template.docx')

def variant(name,identity,birth):
    d=Document(OUT/'team-registration-form-template.docx')
    for table in d.tables:
        for row in list(table.rows):
            text=''.join(_Cell(tc,table).text for tc in row._tr.findall(qn('w:tc')))
            if (not identity and '身份证：' in text) or (not birth and '出生日期：' in text):row._tr.getparent().remove(row._tr)
    d.save(OUT/name)
variant('team-registration-youth-template.docx',False,True)
variant('team-registration-adult-template.docx',True,False)
print('Created generic master, youth and adult Word forms with picture controls and position fields')
